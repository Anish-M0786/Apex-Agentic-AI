from pptx import Presentation
from backend.generators.common import output_path


def create_ppt(filename: str, presentation_title: str, slides: list[dict[str, object]]) -> dict[str, object]:
    path = output_path(filename, '.pptx')
    deck = Presentation()

    # Title slide
    first = deck.slides.add_slide(deck.slide_layouts[0])
    first.shapes.title.text = presentation_title

    # Content slides
    for item in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = str(item.get('title', ''))
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for point in item.get('bullet_points', []) or []:
            frame.add_paragraph().text = str(point)

    deck.save(path)

    # Artifact validation
    if not path.exists():
        raise RuntimeError(f'Presentation was not created at {path}')
    if len(deck.slides) < 2:
        raise RuntimeError('Presentation has no content slides — content generation may have failed')

    return {'success': True, 'file': str(path), 'type': 'pptx'}
