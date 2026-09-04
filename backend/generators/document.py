from docx import Document
from backend.generators.common import output_path


def create_document(filename: str, title: str, paragraphs: list[str]) -> dict[str, object]:
    path = output_path(filename, '.docx')
    doc = Document()
    doc.add_heading(title, 0)
    for paragraph in paragraphs:
        if paragraph.strip():
            doc.add_paragraph(paragraph)
    doc.save(path)

    # Artifact validation
    if not path.exists():
        raise RuntimeError(f'Document was not created at {path}')
    size = path.stat().st_size
    if size < 500:
        raise RuntimeError(f'Document is suspiciously small ({size} bytes) — content may be empty')

    return {'success': True, 'file': str(path), 'type': 'docx'}
